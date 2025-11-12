import logging
from typing import Optional, Dict

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None

from config import SENDER_DATA_XLSX_PATH

_logger = logging.getLogger(__name__)

_cache = {
    "by_id": {},
    "loaded": False,
}


def _normalize_header(name: str) -> str:
    return name.strip().lower()


def _load_data() -> None:
    if _cache["loaded"]:
        return
    _cache["by_id"] = {}

    if not SENDER_DATA_XLSX_PATH:
        _logger.warning("SENDER_DATA_XLSX_PATH is not set; skipping sender data load")
        _cache["loaded"] = True
        return

    if load_workbook is None:
        _logger.warning("openpyxl is not installed; cannot read sender data .xlsx file")
        _cache["loaded"] = True
        return

    try:
        wb = load_workbook(SENDER_DATA_XLSX_PATH, data_only=True, read_only=True)
        ws = wb.active

        # Build header map
        headers = {}
        for cell in next(ws.iter_rows(min_row=1, max_row=1)):
            headers[_normalize_header(str(cell.value or ""))] = cell.column

        # Try to resolve expected columns by keywords (RU)
        def find_col(*candidates):
            for cand in candidates:
                cand_n = _normalize_header(cand)
                for h, col in headers.items():
                    if cand_n == h or cand_n in h:
                        return col
            return None

        col_id = find_col("id", "id отправителя", "id сотрудника", "tg id", "user id", "ид отправителя")
        col_ip = find_col("ип")
        col_email = find_col("почта", "email", "e-mail")
        col_inn = find_col("инн")

        if not col_id:
            _logger.warning("Не найден столбец с ID отправителя в файле: %s", SENDER_DATA_XLSX_PATH)
            _cache["loaded"] = True
            return

        for row in ws.iter_rows(min_row=2):
            try:
                raw_id = row[col_id - 1].value if col_id else None
                if raw_id is None:
                    continue
                # Normalize to string of int if possible
                user_id_str = None
                try:
                    user_id_str = str(int(raw_id))
                except Exception:
                    user_id_str = str(raw_id).strip()

                info = {
                    "ip": (row[col_ip - 1].value if col_ip else None),
                    "email": (row[col_email - 1].value if col_email else None),
                    "inn": (row[col_inn - 1].value if col_inn else None),
                }
                # Normalize text values
                for k, v in list(info.items()):
                    if v is None:
                        info[k] = None
                    else:
                        info[k] = str(v).strip()

                _cache["by_id"][user_id_str] = info
            except Exception as e:
                _logger.debug("Пропущена строка из-за ошибки: %s", e)

        _cache["loaded"] = True
        _logger.info("Загружены данные отправителей из файла: %s (записей: %d)", SENDER_DATA_XLSX_PATH, len(_cache["by_id"]))
    except FileNotFoundError:
        _logger.warning("Файл с данными отправителей не найден: %s", SENDER_DATA_XLSX_PATH)
        _cache["loaded"] = True
    except Exception as e:
        _logger.error("Ошибка при загрузке данных отправителей: %s", e)
        _cache["loaded"] = True


def get_sender_extra_info(user_id: str) -> Optional[Dict[str, Optional[str]]]:
    """Возвращает словарь с ключами ip, email, inn по user_id (строка).
    Если данные не найдены или загрузка невозможна, возвращает None.
    """
    _load_data()
    return _cache["by_id"].get(str(user_id))
